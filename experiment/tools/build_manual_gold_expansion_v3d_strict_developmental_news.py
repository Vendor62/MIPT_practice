#!/usr/bin/env python3.11
"""
build_manual_gold_expansion_v3d_strict_developmental_news

Strict news-only developmental manual gold expansion (v3d).
No LLM, no hybrid metrics, no fulltext in sync_out.
"""

from __future__ import annotations

import argparse
import csv
import importlib.util
import json
import re
import subprocess
import sys
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import numpy as np

_TOOLS = Path(__file__).resolve().parent
if str(_TOOLS) not in sys.path:
    sys.path.insert(0, str(_TOOLS))

_v3_path = _TOOLS / "build_manual_gold_expansion_robustness_v3.py"
_spec_v3 = importlib.util.spec_from_file_location("exp_v3", _v3_path)
exp_v3 = importlib.util.module_from_spec(_spec_v3)
assert _spec_v3 and _spec_v3.loader
_spec_v3.loader.exec_module(exp_v3)

_v3c_path = _TOOLS / "build_manual_gold_expansion_v3c_developmental_stories.py"
_spec_v3c = importlib.util.spec_from_file_location("v3c_build", _v3c_path)
v3c = importlib.util.module_from_spec(_spec_v3c)
assert _spec_v3c and _spec_v3c.loader
_spec_v3c.loader.exec_module(v3c)

MIN_PUBLISHED = datetime(2024, 1, 1, tzinfo=timezone.utc)
TARGET_POSTS = 200
TARGET_MIN = 150
TARGET_MAX = 250
MAX_BLOCKS = 60
MIN_BLOCKS = 30
HIGH_SEM_THRESH = 0.88
MIN_GAP_H = 12.0
MIN_SPAN_H = 48.0
TARGET_SPAN_H = 72.0

EXPANDED_WL = (
    _TOOLS.parent
    / "sync_out/data/prepared/dataset_candidate_mining_whitelist_v1/channel_whitelist_v1.csv"
)

FORBIDDEN_USERNAMES = frozenset(
    {
        "nachemodanah",
        "trendwatching24",
        "rueventjob4at",
        "popsmarketing",
        "breakevens",
        "jobinfood",
        "g00djob4all",
        "worldeventjob",
        "mskeventjob",
        "spbeventjob",
        "talentedpeoples",
        "autonews_youtube",
        "ai_ins",
    }
)
ADS_NAME_RE = re.compile(
    r"(travel|aviasales|tour|ваканс|job|реклам|affiliate|tpx\.lu|level\.tpx|eventjob|marketing)",
    re.I,
)

DEV_KEYWORD_TOPICS: dict[str, tuple[str, ...]] = {
    "dagestan_flood": exp_v3.DAGESTAN_KW,
    "lunar": exp_v3.LUNAR_KW,
    "disaster": (
        "авария",
        "катастроф",
        "пожар",
        "взрыв",
        "чс",
        "эвакуац",
        "мчс",
        "погиб",
    ),
    "legal": ("расследован", "задержан", "приговор", "суд ", "обвинен", "арест"),
    "sanctions": ("санкц", "законопроект", "решение суда", "запрет"),
}

LABEL_COLS = [
    "action",
    "gold_story_id",
    "note",
    "suggested_story_id",
    "candidate_block_id",
    "row_in_block",
    "block_size",
    "weak_storyline_id",
    "candidate_type",
    "temporal_phase",
    "time_gap_from_prev_hours",
    "topic_tags",
    "selection_reason",
    "published_at",
    "channel_name",
    "post_id",
    "text",
]

CAND_FIELDS = [
    "candidate_block_id",
    "source",
    "weak_storyline_id",
    "size",
    "selected_size",
    "channel_count",
    "first_published_at",
    "last_published_at",
    "time_span_hours",
    "temporal_phase_count",
    "median_gap_hours_between_selected_posts",
    "year_month_bucket",
    "story_size_bucket",
    "mean_pairwise_semantic",
    "max_pairwise_semantic",
    "high_semantic_pair_share",
    "lexical_novelty_score",
    "entity_novelty_score",
    "temporal_novelty_score",
    "developmental_score",
    "near_duplicate_score",
    "topic_tags",
    "quality_flags",
    "selection_reason",
    "candidate_type",
    "post_ids",
]


def _exp_root() -> Path:
    return _TOOLS.parent


def _norm_user(name: str) -> str:
    return (name or "").strip().lower().lstrip("@")


def _size_bucket(n: int) -> str:
    return v3c._size_bucket(n)


def _phase_label(hours_from_start: float) -> str:
    return v3c._phase_label(hours_from_start)


def _temporal_novelty(phase_cnt: int, median_gap: float, span_h: float) -> float:
    s = 0.0
    if phase_cnt >= 2:
        s += 0.35
    if phase_cnt >= 3:
        s += 0.2
    if median_gap >= MIN_GAP_H:
        s += 0.25
    if span_h >= TARGET_SPAN_H:
        s += 0.2
    return min(1.0, s)


def _median_gap_hours(pids: list[int], meta: dict[int, dict[str, Any]]) -> float:
    if len(pids) < 2:
        return 0.0
    ts = sorted(meta[p]["published_at"] for p in pids)
    gaps = [(ts[i] - ts[i - 1]).total_seconds() / 3600.0 for i in range(1, len(ts))]
    return float(np.median(gaps)) if gaps else 0.0


def _subsample_temporal_spread(
    pids: list[int],
    meta: dict[int, dict[str, Any]],
    *,
    min_posts: int = 3,
    max_posts: int = 8,
    min_gap_h: float = MIN_GAP_H,
) -> list[int]:
    if len(pids) <= max_posts:
        return pids
    ordered = sorted(pids, key=lambda p: meta[p]["published_at"])
    selected = [ordered[0], ordered[-1]]
    for p in ordered[1:-1]:
        if len(selected) >= max_posts:
            break
        ts = meta[p]["published_at"]
        if all(abs((ts - meta[s]["published_at"]).total_seconds()) / 3600.0 >= min_gap_h for s in selected):
            selected.append(p)
    selected = sorted(set(selected), key=lambda p: meta[p]["published_at"])
    if len(selected) < min_posts:
        return ordered[:max_posts]
    return selected


def _build_strict_whitelist(root: Path) -> tuple[list[dict[str, str]], list[dict[str, str]]]:
    """Strict whitelist from combined_manual_gold_v2 channels only."""
    v2_meta = root / "sync_out/data/prepared/combined_manual_gold_v2/combined_posts_metadata_v2.csv"
    by_cid: dict[str, dict[str, str]] = {}
    for r in csv.DictReader(v2_meta.open(encoding="utf-8")):
        cid = r["channel_id"].strip()
        if cid not in by_cid:
            by_cid[cid] = {
                "channel_id": cid,
                "channel_name": r.get("channel_name", ""),
                "username": "",
                "whitelist_source": "combined_manual_gold_v2",
            }
    wl_ids = sorted(by_cid)
    arr = "ARRAY[" + ",".join("'" + c.replace("'", "''") + "'" for c in wl_ids) + "]::text[]"
    rows = exp_v3._psql_json(
        f"""
        SELECT coalesce(json_agg(row_to_json(t)), '[]'::json)::text FROM (
          SELECT chat_id::text AS channel_id,
                 coalesce(username,'') AS username,
                 coalesce(name,'') AS channel_title
          FROM communities WHERE chat_id::text = ANY({arr})
        ) t;
        """
    )
    for r in rows:
        cid = str(r["channel_id"])
        if cid in by_cid:
            un = (r.get("username") or "").strip()
            by_cid[cid]["username"] = un
            if not by_cid[cid]["channel_name"]:
                by_cid[cid]["channel_name"] = f"@{un}" if un else str(r.get("channel_title") or "")

    strict_rows = [by_cid[c] for c in sorted(by_cid, key=lambda x: by_cid[x]["channel_name"])]
    for row in strict_rows:
        u = _norm_user(row.get("username") or row.get("channel_name", ""))
        if u in FORBIDDEN_USERNAMES or ADS_NAME_RE.search(row.get("channel_name", "")):
            raise RuntimeError(f"forbidden channel in v2 whitelist: {row}")

    excluded: list[dict[str, str]] = []
    if EXPANDED_WL.is_file():
        strict_ids = {r["channel_id"] for r in strict_rows}
        for r in csv.DictReader(EXPANDED_WL.open(encoding="utf-8")):
            cid = r["channel_id"].strip()
            if cid in strict_ids:
                continue
            name = r.get("channel_name", "")
            un = _norm_user(r.get("username") or name)
            reason = []
            if un in FORBIDDEN_USERNAMES or "nachemodanah" in un:
                reason.append("forbidden_username")
            if ADS_NAME_RE.search(name or ""):
                reason.append("ads_travel_job_pattern")
            reason.append("not_in_combined_manual_gold_v2")
            excluded.append(
                {
                    "channel_id": cid,
                    "channel_name": name,
                    "username": un,
                    "exclusion_reason": ";".join(reason),
                }
            )
    return strict_rows, excluded


def _load_exclude(root: Path) -> tuple[set[int], list[str]]:
    warnings: list[str] = []
    exclude: set[int] = set()
    paths = [
        ("combined_v2", root / "sync_out/data/prepared/combined_manual_gold_v2/combined_manual_gold_labels_v2.csv"),
        ("v3", root / "sync_out/data/prepared/manual_gold_expansion_v3/selected_posts_metadata_v3.csv"),
        ("v3b", root / "sync_out/data/prepared/manual_gold_expansion_v3b/selected_posts_metadata_v3b.csv"),
        ("v3c", root / "sync_out/data/prepared/manual_gold_expansion_v3c/selected_posts_metadata_v3c.csv"),
        (
            "v3c_clean",
            root
            / "sync_out/data/prepared/manual_gold_expansion_v3c_clean_audit/v3c_labeled_clean_posts_v1.csv",
        ),
    ]
    for tag, p in paths:
        if not p.is_file():
            warnings.append(f"missing exclusion file: {p.name}")
            continue
        exclude |= {int(r["post_id"]) for r in csv.DictReader(p.open(encoding="utf-8"))}
    return exclude, warnings


def _enrich_v3d(b: dict[str, Any], meta: dict[int, dict[str, Any]], emb: dict[int, np.ndarray], snippets: dict[int, str]) -> dict[str, Any]:
    bb = v3c._enrich_block(b, meta, emb, snippets)
    pids = bb["post_ids"]
    mg = _median_gap_hours(pids, meta)
    span_h = float(bb.get("time_span_hours", 0))
    phase_cnt = int(bb.get("temporal_phase_count", 0))
    bb["median_gap_hours_between_selected_posts"] = round(mg, 2)
    bb["year_month_bucket"] = (bb.get("month_bucket") or "")[:7]
    bb["temporal_novelty_score"] = round(_temporal_novelty(phase_cnt, mg, span_h), 4)
    bb["selected_size"] = len(pids)
    return bb


def _apply_subsample_and_reenrich(
    b: dict[str, Any],
    meta: dict[int, dict[str, Any]],
    emb: dict[int, np.ndarray],
    snippets: dict[int, str],
) -> dict[str, Any]:
    raw = dict(b)
    subs = _subsample_temporal_spread(b["post_ids"], meta)
    if subs != b["post_ids"]:
        raw["post_ids"] = subs
        raw["selection_reason"] = (b.get("selection_reason") or "") + ";temporal_subsample"
    return _enrich_v3d(raw, meta, emb, snippets)


def _passes_developmental_quality(b: dict[str, Any]) -> bool:
    if b.get("candidate_type") == "reject":
        return False
    span = float(b.get("time_span_hours", 0))
    if span < MIN_SPAN_H:
        return False
    if span <= 12 and float(b.get("high_semantic_pair_share", 0)) >= 0.65:
        return False
    if float(b.get("near_duplicate_score", 0)) >= 0.55:
        return False
    if float(b.get("developmental_score", 0)) < 0.48:
        return False
    if int(b.get("temporal_phase_count", 0)) < 2:
        return False
    if float(b.get("lexical_novelty_score", 0)) < 0.04 and float(b.get("mean_pairwise_semantic", 0)) > 0.9:
        return False
    return True


def _mine_neo_developmental_v3d(
    exclude: set[int],
    wl: set[str],
    date_from: datetime,
    date_to: datetime,
) -> list[dict[str, Any]]:
    blocks: list[dict[str, Any]] = []
    seen: set[str] = set()
    for skip in range(0, 20000, 400):
        neo_blocks = exp_v3._neo_storyline_blocks(skip, 400)
        if not neo_blocks:
            break
        all_pids: list[int] = []
        for _, pids in neo_blocks:
            all_pids.extend(pids)
        meta = exp_v3._fetch_posts_meta(list(set(all_pids)), wl)
        need = [p for p in all_pids if p in meta and p not in exclude]
        emb = exp_v3._load_embeddings(need)
        snippets = v3c._fetch_snippets(need)

        for sid, pids in neo_blocks:
            filt = [p for p in pids if p in meta and p not in exclude]
            if len(filt) < 3 or len(filt) > 12:
                continue
            filt.sort(key=lambda p: meta[p]["published_at"])
            ts_list = [meta[p]["published_at"] for p in filt]
            if min(ts_list) < MIN_PUBLISHED or max(ts_list) > date_to or min(ts_list) > date_to:
                continue
            if any(meta[p]["text_len"] < 120 for p in filt):
                continue
            span_h = (max(ts_list) - min(ts_list)).total_seconds() / 3600.0
            if span_h < MIN_SPAN_H:
                continue
            bid = f"neo_dev_{sid}_{len(filt)}"
            if bid in seen:
                continue
            seen.add(bid)
            raw = {
                "candidate_block_id": bid,
                "source": "neo4j_storyline",
                "weak_storyline_id": sid,
                "post_ids": filt,
                "graph_cohesion_score": 0.3,
                "entity_overlap_proxy": "",
                "selection_reason": "neo_storyline_strict_developmental",
            }
            bb = _apply_subsample_and_reenrich(raw, meta, emb, snippets)
            if _passes_developmental_quality(bb):
                blocks.append(bb)
    return blocks


def _mine_keyword_developmental(
    exclude: set[int],
    wl_list: list[str],
    wl: set[str],
    date_from: datetime,
    date_to: datetime,
) -> list[dict[str, Any]]:
    blocks: list[dict[str, Any]] = []
    seen: set[frozenset[int]] = set()
    for topic, kws in DEV_KEYWORD_TOPICS.items():
        for b in v3c._mine_topic_window_developmental(exclude, wl_list, wl, date_from, date_to, topic, kws):
            key = frozenset(b["post_ids"])
            if key in seen:
                continue
            seen.add(key)
            meta = exp_v3._fetch_posts_meta(b["post_ids"], wl)
            emb = exp_v3._load_embeddings(b["post_ids"])
            snippets = v3c._fetch_snippets(b["post_ids"])
            bb = _apply_subsample_and_reenrich(b, meta, emb, snippets)
            if _passes_developmental_quality(bb):
                blocks.append(bb)
    return blocks


def _mine_temporal_entity_windows(
    exclude: set[int],
    wl_list: list[str],
    wl: set[str],
    date_from: datetime,
    date_to: datetime,
) -> list[dict[str, Any]]:
    """2-week windows with span>=48h and size>=3 on strict channels."""
    ch_arr = "ARRAY[" + ",".join("'" + c.replace("'", "''") + "'" for c in wl_list) + "]::text[]"
    excl = ""
    if exclude:
        excl = " AND id NOT IN (" + ",".join(str(x) for x in sorted(exclude)) + ")"
    sql = f"""
    SELECT coalesce(json_agg(row_to_json(t)), '[]'::json)::text FROM (
      SELECT id AS post_id, timestamp AS published_at, channel_id
      FROM public.posts
      WHERE timestamp >= '{date_from.isoformat()}' AND timestamp <= '{date_to.isoformat()}'
        AND channel_id = ANY({ch_arr}) AND length(coalesce(content,'')) >= 200 {excl}
      ORDER BY timestamp
      LIMIT 12000
    ) t;
    """
    rows = exp_v3._psql_json(sql)
    by_month: dict[str, list[int]] = defaultdict(list)
    for r in rows:
        ts = exp_v3._parse_ts(str(r["published_at"]))
        if not ts:
            continue
        by_month[ts.strftime("%Y-%m")].append(int(r["post_id"]))
    blocks: list[dict[str, Any]] = []
    seen: set[frozenset[int]] = set()
    for month, pids in by_month.items():
        if len(pids) < 8:
            continue
        meta = exp_v3._fetch_posts_meta(pids[:400], wl)
        ordered = sorted(meta.keys(), key=lambda p: meta[p]["published_at"])
        win: list[int] = []
        for pid in ordered:
            win.append(pid)
            if len(win) < 3:
                continue
            ts_w = [meta[p]["published_at"] for p in win]
            span_h = (max(ts_w) - min(ts_w)).total_seconds() / 3600.0
            if span_h > 14 * 24:
                win = win[-3:]
                ts_w = [meta[p]["published_at"] for p in win]
                span_h = (max(ts_w) - min(ts_w)).total_seconds() / 3600.0
            if span_h >= MIN_SPAN_H and len(win) >= 3:
                key = frozenset(win)
                if key not in seen:
                    seen.add(key)
                    raw = {
                        "candidate_block_id": f"tw_{month}_{win[0]}",
                        "source": "temporal_window",
                        "weak_storyline_id": "",
                        "post_ids": list(win),
                        "selection_reason": "temporal_window_month",
                    }
                    emb = exp_v3._load_embeddings(win)
                    snippets = v3c._fetch_snippets(win)
                    bb = _apply_subsample_and_reenrich(raw, meta, emb, snippets)
                    if _passes_developmental_quality(bb):
                        blocks.append(bb)
            if len(win) > 10:
                win = win[-5:]
    return blocks


def _mine_singleton_controls(
    exclude: set[int],
    wl_list: list[str],
    date_from: datetime,
    date_to: datetime,
    limit: int = 250,
) -> list[dict[str, Any]]:
    return v3c._mine_singleton_controls(exclude, wl_list, date_from, date_to, limit)


def _mine_near_dup(pool: list[dict[str, Any]], limit: int = 10) -> list[dict[str, Any]]:
    return v3c._mine_near_duplicate_diagnostic(pool, limit)


def _tags(b: dict[str, Any]) -> set[str]:
    return set((b.get("topic_tags") or "").split(","))


def _select_v3d(
    developmental: list[dict[str, Any]],
    singletons: list[dict[str, Any]],
    near_dup: list[dict[str, Any]],
    hard_neg: list[dict[str, Any]],
    target: int = TARGET_POSTS,
) -> list[dict[str, Any]]:
    used: set[int] = set()
    chosen: list[dict[str, Any]] = []
    posts_by_type: Counter[str] = Counter()
    months_posts: Counter[str] = Counter()
    quotas = {
        "developmental": int(0.72 * target),
        "singleton_control": int(0.18 * target),
        "near_duplicate_diagnostic": int(0.08 * target),
        "hard_negative": int(0.06 * target),
    }

    def month_cap_ok(month: str, add: int) -> bool:
        total = sum(months_posts.values()) + add
        if months_posts[month] + add > max(45, int(0.28 * total)):
            return False
        if month in ("2026-04", "2026-05") and months_posts[month] + add > int(0.32 * total):
            return False
        return True

    def add_block(b: dict[str, Any], ctype: str, reason: str) -> bool:
        pids = [p for p in b["post_ids"] if p not in used]
        if not pids:
            return False
        month = (b.get("month_bucket") or "")[:7]
        if not month_cap_ok(month, len(pids)):
            return False
        if sum(posts_by_type.values()) + len(pids) > TARGET_MAX:
            return False
        if len(chosen) >= MAX_BLOCKS:
            return False
        if posts_by_type[ctype] + len(pids) > quotas[ctype] + 6:
            return False
        bb = dict(b)
        bb["post_ids"] = pids
        bb["size"] = len(pids)
        bb["selected_size"] = len(pids)
        bb["story_size_bucket"] = _size_bucket(len(pids))
        bb["candidate_type"] = ctype
        bb["selection_reason_v3d"] = reason
        chosen.append(bb)
        used.update(pids)
        posts_by_type[ctype] += len(pids)
        months_posts[month] += len(pids)
        return True

    dev_sorted = sorted(
        developmental,
        key=lambda x: (
            -float(x.get("temporal_novelty_score", 0)),
            -float(x.get("developmental_score", 0)),
            -float(x.get("time_span_hours", 0)),
            -int(x.get("temporal_phase_count", 0)),
            -float(x.get("lexical_novelty_score", 0)),
            float(x.get("high_semantic_pair_share", 1)),
        ),
    )
    topics_seen: Counter[str] = Counter()

    for b in near_dup:
        if posts_by_type["near_duplicate_diagnostic"] >= quotas["near_duplicate_diagnostic"]:
            break
        add_block(b, "near_duplicate_diagnostic", "near_duplicate_diagnostic")

    for b in hard_neg:
        if posts_by_type["hard_negative"] >= quotas["hard_negative"]:
            break
        add_block(b, "hard_negative", "hard_negative")

    for b in dev_sorted:
        if posts_by_type["developmental"] >= quotas["developmental"]:
            break
        topic_key = next((t for t in _tags(b) if t != "other"), "other")
        if topics_seen[topic_key] >= 5:
            continue
        if add_block(b, "developmental", "developmental_strict"):
            topics_seen[topic_key] += 1

    for b in sorted(singletons, key=lambda x: x.get("month_bucket", "")):
        if posts_by_type["singleton_control"] >= quotas["singleton_control"]:
            break
        add_block(b, "singleton_control", "singleton_control")

    for b in dev_sorted:
        if sum(posts_by_type.values()) >= TARGET_MIN:
            break
        add_block(b, "developmental", "developmental_fill")

    return chosen


def _write_labeling_sheet(rows: list[dict[str, str]], csv_path: Path, xlsx_path: Path) -> None:
    csv_path.parent.mkdir(parents=True, exist_ok=True)
    with csv_path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=LABEL_COLS)
        w.writeheader()
        for r in rows:
            w.writerow({k: r.get(k, "") for k in LABEL_COLS})

    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    type_fills = {
        "developmental": PatternFill("solid", fgColor="C6EFCE"),
        "singleton_control": PatternFill("solid", fgColor="DDEBF7"),
        "near_duplicate_diagnostic": PatternFill("solid", fgColor="FCE4D6"),
        "hard_negative": PatternFill("solid", fgColor="E2EFDA"),
    }
    wb = openpyxl.Workbook()
    ins = wb.active
    ins.title = "instructions"
    ins["A1"] = "Manual gold expansion v3d — strict developmental news"
    for i, t in enumerate(
        [
            "gold_story_id — главное поле (оставьте пустым до разметки).",
            "suggested_story_id — подсказка EXP3D_BLOCK_xxxx, не gold.",
            "Один развивающийся сюжет в блоке → один gold_story_id на все строки.",
            "Несколько сюжетов в блоке → разные gold_story_id.",
            "singleton_control / hard_negative — обычно отдельный gold_story_id.",
            "action=drop для рекламы, перепечатки без развития, нерелевантного поста.",
            "note — для спорных split/merge.",
            "temporal_phase и time_gap_from_prev_hours — эвристика для фаз сюжета.",
        ],
        start=3,
    ):
        ins[f"A{i}"] = t

    ws = wb.create_sheet("review")
    header_fill = PatternFill("solid", fgColor="D9E1F2")
    for j, col in enumerate(LABEL_COLS, start=1):
        c = ws.cell(row=1, column=j, value=col)
        c.font = Font(bold=True)
        c.fill = header_fill
    prev_block = None
    for i, r in enumerate(rows, start=2):
        is_first = r["candidate_block_id"] != prev_block
        prev_block = r["candidate_block_id"]
        ctype = r.get("candidate_type", "")
        row_fill = type_fills.get(ctype)
        for j, col in enumerate(LABEL_COLS, start=1):
            cell = ws.cell(row=i, column=j, value=r.get(col, ""))
            if col == "text":
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            if row_fill and (is_first or col in ("candidate_type", "candidate_block_id", "suggested_story_id")):
                cell.fill = row_fill
        ws.row_dimensions[i].height = 72 if r.get("text") else 22
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(LABEL_COLS))}{len(rows)+1}"
    widths = {
        "A": 8,
        "B": 14,
        "C": 16,
        "D": 18,
        "E": 20,
        "F": 5,
        "G": 7,
        "H": 14,
        "I": 22,
        "J": 18,
        "K": 14,
        "L": 16,
        "M": 22,
        "N": 20,
        "O": 12,
        "P": 10,
        "Q": 55,
    }
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    wb.save(xlsx_path)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--target-posts", type=int, default=TARGET_POSTS)
    args = ap.parse_args()

    root = _exp_root()
    out_dir = root / "sync_out/data/prepared/manual_gold_expansion_v3d"
    reports = root / "sync_out/reports"
    sync_in = root / "sync_in"
    for p in (out_dir, reports, sync_in):
        p.mkdir(parents=True, exist_ok=True)

    strict_rows, excluded_expanded = _build_strict_whitelist(root)
    wl_path = out_dir / "strict_news_whitelist_v3d.csv"
    with wl_path.open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["channel_id", "channel_name", "username", "whitelist_source"])
        w.writeheader()
        w.writerows(strict_rows)
    wl_ids = {r["channel_id"] for r in strict_rows}
    wl_list = sorted(wl_ids)

    exclude, excl_warnings = _load_exclude(root)
    max_ts = exp_v3._parse_ts(exp_v3._psql_scalar("SELECT max(timestamp)::text FROM public.posts WHERE timestamp >= '2024-01-01';"))
    assert max_ts
    date_from = MIN_PUBLISHED
    date_to = max_ts

    neo_pool = _mine_neo_developmental_v3d(exclude, wl_ids, date_from, date_to)
    kw_pool = _mine_keyword_developmental(exclude, wl_list, wl_ids, date_from, date_to)
    tw_pool = _mine_temporal_entity_windows(exclude, wl_list, wl_ids, date_from, date_to)

    merged: list[dict[str, Any]] = []
    seen_p: set[frozenset[int]] = set()
    for b in neo_pool + kw_pool + tw_pool:
        key = frozenset(b["post_ids"])
        if key in seen_p:
            continue
        seen_p.add(key)
        merged.append(b)

    developmental = [b for b in merged if b.get("candidate_type") != "reject"]
    near_pool = [b for b in merged if float(b.get("near_duplicate_score", 0)) >= 0.5]
    singletons = _mine_singleton_controls(exclude, wl_list, date_from, date_to)
    hard_neg = v3c._mine_hard_negatives(developmental, merged, exclude, wl_ids, max_blocks=6)

    all_cand = merged + singletons
    with (out_dir / "candidate_blocks_v3d.csv").open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=CAND_FIELDS, extrasaction="ignore")
        w.writeheader()
        for b in all_cand:
            row = {k: b.get(k, "") for k in CAND_FIELDS}
            row["post_ids"] = ";".join(str(p) for p in b["post_ids"])
            w.writerow(row)

    dev_spans = [float(b["time_span_hours"]) for b in developmental if int(b.get("size", 0)) >= 3]
    reports.joinpath("manual_gold_expansion_candidate_inventory_v3d.md").write_text(
        "# Candidate inventory v3d (strict news)\n\n"
        f"- Strict whitelist channels: **{len(strict_rows)}** (source: combined_manual_gold_v2)\n"
        f"- Expanded whitelist excluded channels: **{len(excluded_expanded)}**\n"
        f"- Developmental candidates (quality pass): **{len(developmental)}**\n"
        f"- Neo developmental blocks: **{len(neo_pool)}**\n"
        f"- Keyword developmental: **{len(kw_pool)}**\n"
        f"- Temporal window: **{len(tw_pool)}**\n"
        f"- Median span (dev, h): **{float(np.median(dev_spans)) if dev_spans else 0:.1f}**\n"
        f"- Date window: {date_from.date()} … {date_to.date()}\n"
        + ("\n".join(f"- WARN: {w}" for w in excl_warnings) if excl_warnings else "")
        + "\n",
        encoding="utf-8",
    )

    chosen = _select_v3d(
        developmental,
        singletons,
        _mine_near_dup(near_pool, limit=10),
        hard_neg,
        target=args.target_posts,
    )

    label_rows: list[dict[str, str]] = []
    meta_rows: list[dict[str, str]] = []
    block_rows: list[dict[str, str]] = []
    pids_all = [p for b in chosen for p in b["post_ids"]]
    arr = "ARRAY[" + ",".join(str(x) for x in pids_all) + "]::bigint[]"
    pg_rows = exp_v3._psql_json(
        f"""
        SELECT coalesce(json_agg(row_to_json(t) ORDER BY t.post_id), '[]'::json)::text FROM (
          SELECT p.id AS post_id, p.timestamp AS published_at, p.channel_id,
                 coalesce(c.username, '') AS username, coalesce(c.name, '') AS channel_title,
                 coalesce(p.content, '') AS content
          FROM public.posts p
          LEFT JOIN communities c ON c.chat_id::text = p.channel_id
          WHERE p.id = ANY({arr})
        ) t;
        """
    )
    pg_map = {int(r["post_id"]): r for r in pg_rows}

    months: Counter[str] = Counter()
    years: Counter[str] = Counter()
    old2023 = 0
    min_dt = max_dt = None
    posts_by_type: Counter[str] = Counter()
    dev_spans_sel: list[float] = []
    dev_gaps: list[float] = []
    phase_counts: list[int] = []

    for bi, b in enumerate(chosen, start=1):
        sug = f"EXP3D_BLOCK_{bi:04d}"
        ctype = b.get("candidate_type", "developmental")
        block_rows.append(
            {
                "candidate_block_id": b["candidate_block_id"],
                "suggested_story_id": sug,
                "candidate_type": ctype,
                "n_posts": len(b["post_ids"]),
                "time_span_hours": b.get("time_span_hours", ""),
                "median_gap_hours": b.get("median_gap_hours_between_selected_posts", ""),
                "temporal_phase_count": b.get("temporal_phase_count", ""),
                "developmental_score": b.get("developmental_score", ""),
                "selection_reason_v3d": b.get("selection_reason_v3d", ""),
            }
        )
        if ctype == "developmental":
            dev_spans_sel.append(float(b.get("time_span_hours") or 0))
            dev_gaps.append(float(b.get("median_gap_hours_between_selected_posts") or 0))
            phase_counts.append(int(b.get("temporal_phase_count") or 0))

        sorted_pids = sorted(
            b["post_ids"],
            key=lambda p: exp_v3._parse_ts(str(pg_map[p]["published_at"])) or MIN_PUBLISHED,
        )
        prev_ts = None
        for ri, pid in enumerate(sorted_pids):
            pg = pg_map[pid]
            ts = exp_v3._parse_ts(str(pg["published_at"]))
            if not ts:
                continue
            gap_h = ""
            if prev_ts is not None:
                gap_h = f"{(ts - prev_ts).total_seconds() / 3600.0:.2f}"
            prev_ts = ts
            t0 = exp_v3._parse_ts(str(pg_map[sorted_pids[0]]["published_at"])) or ts
            hours = (ts - t0).total_seconds() / 3600.0
            phase = _phase_label(hours) if ctype == "developmental" else "singleton"
            if ts < MIN_PUBLISHED:
                old2023 += 1
            min_dt = ts if min_dt is None else min(min_dt, ts)
            max_dt = ts if max_dt is None else max(max_dt, ts)
            months[ts.strftime("%Y-%m")] += 1
            years[ts.strftime("%Y")] += 1
            posts_by_type[ctype] += 1
            uname = (pg.get("username") or "").strip()
            ch_name = uname if uname.startswith("@") else (f"@{uname}" if uname else str(pg.get("channel_title") or ""))
            text = exp_v3._normalize_text(str(pg.get("content") or ""))
            tags = b.get("topic_tags") or ",".join(exp_v3._topic_tags_from_text(text))
            meta_rows.append(
                {
                    "post_id": pid,
                    "published_at": ts.isoformat(),
                    "channel_id": str(pg.get("channel_id") or ""),
                    "channel_name": ch_name,
                    "candidate_block_id": b["candidate_block_id"],
                    "suggested_story_id": sug,
                    "candidate_type": ctype,
                    "block_size": len(b["post_ids"]),
                    "story_size_bucket": b.get("story_size_bucket", ""),
                    "topic_tags": tags,
                }
            )
            label_rows.append(
                {
                    "action": "",
                    "gold_story_id": "",
                    "note": "",
                    "suggested_story_id": sug,
                    "candidate_block_id": b["candidate_block_id"],
                    "row_in_block": ri + 1,
                    "block_size": len(b["post_ids"]),
                    "weak_storyline_id": b.get("weak_storyline_id", ""),
                    "candidate_type": ctype,
                    "temporal_phase": phase,
                    "time_gap_from_prev_hours": gap_h,
                    "topic_tags": tags,
                    "selection_reason": b.get("selection_reason_v3d", b.get("selection_reason", "")),
                    "published_at": ts.strftime("%Y-%m-%d %H:%M:%S%z"),
                    "channel_name": ch_name,
                    "post_id": pid,
                    "text": text,
                }
            )

    total = len(label_rows)
    bucket_posts = Counter(_size_bucket(len(b["post_ids"])) for b in chosen)

    manifest = {
        "dataset": "manual_gold_expansion_v3d",
        "strict_whitelist_source": "combined_manual_gold_v2",
        "strict_whitelist_path": str(wl_path.relative_to(root)),
        "strict_whitelist_channel_count": len(strict_rows),
        "expanded_whitelist_excluded_count": len(excluded_expanded),
        "target_posts": args.target_posts,
        "selected_posts": total,
        "selected_blocks": len(chosen),
        "exclusion_warnings": excl_warnings,
        "overlap_combined_manual_gold_v2": 0,
        "candidate_type_post_counts": dict(posts_by_type),
        "developmental_post_rate": round(posts_by_type["developmental"] / max(total, 1), 4),
        "singleton_control_post_rate": round(posts_by_type["singleton_control"] / max(total, 1), 4),
        "near_duplicate_diagnostic_post_rate": round(
            posts_by_type["near_duplicate_diagnostic"] / max(total, 1), 4
        ),
        "hard_negative_post_rate": round(posts_by_type["hard_negative"] / max(total, 1), 4),
        "story_size_bucket_distribution": dict(bucket_posts),
        "developmental_median_time_span_hours": float(np.median(dev_spans_sel)) if dev_spans_sel else 0.0,
        "developmental_mean_time_span_hours": float(np.mean(dev_spans_sel)) if dev_spans_sel else 0.0,
        "developmental_median_gap_hours": float(np.median(dev_gaps)) if dev_gaps else 0.0,
        "developmental_blocks_phase_ge2_share": round(
            sum(1 for c in phase_counts if c >= 2) / max(len(phase_counts), 1), 4
        ),
        "developmental_blocks_phase_ge3_share": round(
            sum(1 for c in phase_counts if c >= 3) / max(len(phase_counts), 1), 4
        ),
        "published_at_min": min_dt.isoformat() if min_dt else "",
        "published_at_max": max_dt.isoformat() if max_dt else "",
        "month_distribution": dict(months),
        "year_distribution": dict(years),
        "2023_or_older_count": old2023,
        "forbidden_channels_present": False,
        "candidate_pool_developmental": len(developmental),
        "date_window": {"from": date_from.isoformat(), "to": date_to.isoformat()},
    }

    (out_dir / "selection_manifest_v3d.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    if block_rows:
        with (out_dir / "selected_blocks_v3d.csv").open("w", encoding="utf-8", newline="") as f:
            w = csv.DictWriter(f, fieldnames=list(block_rows[0].keys()))
            w.writeheader()
            w.writerows(block_rows)
    if meta_rows:
        with (out_dir / "selected_posts_metadata_v3d.csv").open("w", encoding="utf-8", newline="") as f:
            w = csv.DictWriter(f, fieldnames=list(meta_rows[0].keys()))
            w.writeheader()
            w.writerows(meta_rows)

    excl_report = reports / "manual_gold_expansion_v3d_expanded_whitelist_exclusions.md"
    top_excl = sorted(excluded_expanded, key=lambda x: x.get("channel_name", ""))[:25]
    excl_report.write_text(
        "# Expanded whitelist exclusions (v3d)\n\n"
        f"- Strict channels: **{len(strict_rows)}**\n"
        f"- Excluded from expanded 58-channel list: **{len(excluded_expanded)}**\n\n"
        "## Sample excluded channels\n\n"
        + "\n".join(f"- {r['channel_name']} ({r['channel_id']}): {r['exclusion_reason']}" for r in top_excl)
        + ("\n- ..." if len(excluded_expanded) > 25 else "")
        + "\n",
        encoding="utf-8",
    )

    reports.joinpath("manual_gold_expansion_selection_v3d.md").write_text(
        "# Selection v3d (strict developmental news)\n\n"
        f"- Posts: **{total}**, blocks: **{len(chosen)}**\n"
        f"- Types: `{dict(posts_by_type)}`\n"
        f"- Dev rate: **{manifest['developmental_post_rate']:.1%}**\n"
        f"- Median dev span: **{manifest['developmental_median_time_span_hours']:.1f}** h\n"
        f"- Median gap: **{manifest['developmental_median_gap_hours']:.1f}** h\n"
        f"- Years: `{dict(years)}`\n",
        encoding="utf-8",
    )

    csv_label = sync_in / "grouped_manual_gold_expansion_v3d_strict_developmental_labeling.csv"
    xlsx_label = sync_in / "grouped_manual_gold_expansion_v3d_strict_developmental_labeling.xlsx"
    _write_labeling_sheet(label_rows, csv_label, xlsx_label)

    print(json.dumps(manifest, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
