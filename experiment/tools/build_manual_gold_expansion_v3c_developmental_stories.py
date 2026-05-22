#!/usr/bin/env python3.11
"""
build_manual_gold_expansion_v3c_developmental_stories

Developmental-story expansion set for manual gold labeling.
Focus: evolving storylines, not near-duplicate repost clusters.
"""

from __future__ import annotations

import argparse
import csv
import importlib.util
import json
import math
import re
import sys
from collections import Counter, defaultdict
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any

import numpy as np

_TOOLS = Path(__file__).resolve().parent
if str(_TOOLS) not in sys.path:
    sys.path.insert(0, str(_TOOLS))

_v3_path = _TOOLS / "build_manual_gold_expansion_robustness_v3.py"
_spec = importlib.util.spec_from_file_location("exp_v3", _v3_path)
assert _spec and _spec.loader
exp_v3 = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(exp_v3)

MIN_PUBLISHED = datetime(2024, 1, 1, tzinfo=timezone.utc)
PRIMARY_START = datetime(2025, 1, 1, tzinfo=timezone.utc)
TARGET_POSTS = 270
TARGET_MIN = 220
TARGET_MAX = 320
MAX_BLOCKS = 90
MIN_BLOCKS = 50
TRUMP_IRAN_CAP = 0.15
NEAR_DUP_CAP = 0.15
DEV_POST_MIN = 0.55
SING_POST_MAX = 0.35
SING_POST_MIN = 0.15
HIGH_SEM_THRESH = 0.88

LABEL_COLS = [
    "action",
    "gold_story_id",
    "note",
    "suggested_story_id",
    "candidate_block_id",
    "row_in_block",
    "block_size",
    "weak_storyline_id",
    "candidate_type",
    "temporal_phase",
    "topic_tags",
    "selection_reason",
    "published_at",
    "channel_name",
    "post_id",
    "text",
]

TOKEN_RE = re.compile(r"[a-zа-яё0-9]{4,}", re.I)
NUM_RE = re.compile(r"\b\d{1,6}\b")

CAND_FIELDS = [
    "candidate_block_id",
    "source",
    "weak_storyline_id",
    "size",
    "channel_count",
    "first_published_at",
    "last_published_at",
    "time_span_hours",
    "temporal_phase_count",
    "month_bucket",
    "story_size_bucket",
    "mean_pairwise_semantic",
    "max_pairwise_semantic",
    "high_semantic_pair_share",
    "lexical_novelty_score",
    "entity_novelty_score",
    "graph_cohesion_score",
    "entity_overlap_proxy",
    "developmental_score",
    "near_duplicate_score",
    "topic_tags",
    "quality_flags",
    "selection_reason",
    "candidate_type",
    "post_ids",
]


def _exp_root() -> Path:
    return _TOOLS.parent


def _size_bucket(n: int) -> str:
    return exp_v3._size_bucket(n)


def _span_bucket(hours: float) -> str:
    if hours <= 6:
        return "same-window-0-6h"
    if hours <= 24:
        return "same-day-6-24h"
    if hours <= 72:
        return "1-3d"
    return "3d+"


def _phase_label(hours_from_start: float) -> str:
    if hours_from_start <= 6:
        return "phase_1_initial"
    if hours_from_start <= 24:
        return "phase_2_update"
    if hours_from_start <= 72:
        return "phase_3_followup"
    return "phase_4_later"


def _token_set(text: str) -> set[str]:
    return set(TOKEN_RE.findall((text or "").lower()))


def _num_set(text: str) -> set[str]:
    return set(NUM_RE.findall(text or ""))


def _pairwise_sem_extended(pids: list[int], emb: dict[int, np.ndarray]) -> tuple[float, float, float, float]:
    vecs = [emb[p] for p in pids if p in emb]
    if len(vecs) < 2:
        return 0.0, 0.0, 0.0, 0.0
    mat = np.stack(vecs, axis=0)
    mat = mat / (np.linalg.norm(mat, axis=1, keepdims=True) + 1e-12)
    sims: list[float] = []
    for i in range(len(vecs)):
        for j in range(i + 1, len(vecs)):
            sims.append(float(np.dot(mat[i], mat[j])))
    if not sims:
        return 0.0, 0.0, 0.0, 0.0
    high = sum(1 for s in sims if s >= HIGH_SEM_THRESH) / len(sims)
    return float(np.mean(sims)), float(np.min(sims)), float(np.max(sims)), high


def _temporal_phase_count(ts_list: list[datetime]) -> int:
    if len(ts_list) < 2:
        return 1
    t0 = min(ts_list)
    phases: set[str] = set()
    for ts in ts_list:
        h = (ts - t0).total_seconds() / 3600.0
        phases.add(_phase_label(h))
    return len(phases)


def _lexical_novelty(texts: list[str]) -> tuple[float, float]:
    if not texts:
        return 0.0, 0.0
    sets = [_token_set(t) for t in texts]
    nums = [_num_set(t) for t in texts]
    base = sets[0] | nums[0]
    if not base:
        return 0.0, 0.0
    novelties: list[float] = []
    entity_jaccards: list[float] = []
    for i in range(1, len(texts)):
        cur = sets[i] | nums[i]
        new_toks = cur - base
        novelties.append(len(new_toks) / max(len(cur), 1))
        inter = len(sets[i] & sets[0])
        union = len(sets[i] | sets[0]) or 1
        entity_jaccards.append(1.0 - inter / union)
        base |= cur
    if not novelties:
        return 0.0, 0.0
    return float(np.mean(novelties)), float(np.mean(entity_jaccards))


def _fetch_snippets(pids: list[int]) -> dict[int, str]:
    if not pids:
        return {}
    out: dict[int, str] = {}
    chunk = 200
    for i in range(0, len(pids), chunk):
        sub = pids[i : i + chunk]
        arr = "ARRAY[" + ",".join(str(x) for x in sub) + "]::bigint[]"
        rows = exp_v3._psql_json(
            f"SELECT coalesce(json_agg(row_to_json(t)), '[]'::json)::text FROM ("
            f"SELECT id AS post_id, left(content, 5000) AS snippet "
            f"FROM public.posts WHERE id = ANY({arr})) t;"
        )
        for r in rows:
            out[int(r["post_id"])] = r.get("snippet") or ""
    return out


def _enrich_block(
    b: dict[str, Any],
    meta: dict[int, dict[str, Any]],
    emb: dict[int, np.ndarray],
    snippets: dict[int, str],
) -> dict[str, Any]:
    pids = sorted(b["post_ids"], key=lambda p: meta[p]["published_at"])
    ts_list = [meta[p]["published_at"] for p in pids]
    texts = [snippets.get(p, "") for p in pids]
    span_h = (max(ts_list) - min(ts_list)).total_seconds() / 3600.0 if len(ts_list) >= 2 else 0.0
    mean_s, _min_s, max_s, high_share = _pairwise_sem_extended(pids, emb)
    phase_cnt = _temporal_phase_count(ts_list)
    lex_nov, ent_nov = _lexical_novelty(texts)
    chans = {meta[p]["channel_id"] for p in pids}
    tags: set[str] = set()
    for t in texts:
        tags.update(exp_v3._topic_tags_from_text(t))
    tags_s = ",".join(sorted(tags)) if tags else "other"

    near_dup = 0.0
    if span_h <= 6 and mean_s >= 0.88 and high_share >= 0.7:
        near_dup += 0.5
    if span_h <= 3 and max_s >= 0.92:
        near_dup += 0.3
    if phase_cnt <= 1 and span_h <= 12:
        near_dup += 0.2
    if lex_nov < 0.05 and mean_s > 0.85:
        near_dup += 0.25
    near_dup = min(1.0, near_dup)

    dev = 0.0
    if span_h >= 12:
        dev += 0.15
    if span_h >= 24:
        dev += 0.2
    if len(pids) >= 3:
        dev += 0.1
    if phase_cnt >= 2:
        dev += 0.25
    if phase_cnt >= 3:
        dev += 0.1
    if 0.35 <= mean_s <= 0.88:
        dev += 0.15
    if high_share < 0.5:
        dev += 0.1
    if lex_nov >= 0.08:
        dev += 0.15
    if ent_nov >= 0.15:
        dev += 0.1
    if len(chans) >= 2:
        dev += 0.08
    if span_h > 6 and phase_cnt >= 2:
        dev += 0.1
    dev -= near_dup * 0.6
    if "lunar" in tags or "dagestan_flood" in tags:
        dev += 0.08
    dev = max(0.0, min(1.0, dev))

    qf: list[str] = []
    if near_dup >= 0.55:
        qf.append("near_duplicate_risk")
    if dev >= 0.45 and near_dup < 0.55:
        qf.append("developmental_candidate")
    if span_h <= 3 and len(pids) >= 3:
        qf.append("burst_window")

    ctype = "developmental"
    if near_dup >= 0.55:
        ctype = "near_duplicate_diagnostic"
    elif dev < 0.35:
        ctype = "reject"

    bb = dict(b)
    bb.update(
        {
            "post_ids": pids,
            "size": len(pids),
            "story_size_bucket": _size_bucket(len(pids)),
            "channel_count": len(chans),
            "time_span_hours": span_h,
            "time_span_bucket": _span_bucket(span_h),
            "temporal_phase_count": phase_cnt,
            "first_published_at": min(ts_list).isoformat(),
            "last_published_at": max(ts_list).isoformat(),
            "month_bucket": min(ts_list).strftime("%Y-%m"),
            "mean_pairwise_semantic": mean_s,
            "max_pairwise_semantic": max_s,
            "high_semantic_pair_share": high_share,
            "lexical_novelty_score": lex_nov,
            "entity_novelty_score": ent_nov,
            "topic_tags": tags_s,
            "developmental_score": dev,
            "near_duplicate_score": near_dup,
            "candidate_type": ctype,
            "quality_flags": ";".join(qf) if qf else "",
        }
    )
    return bb


def _mine_neo_developmental(
    exclude: set[int],
    wl: set[str],
    date_from: datetime,
    date_to: datetime,
) -> list[dict[str, Any]]:
    blocks: list[dict[str, Any]] = []
    seen: set[str] = set()
    for skip in range(0, 16000, 400):
        neo_blocks = exp_v3._neo_storyline_blocks(skip, 400)
        if not neo_blocks:
            break
        all_pids: list[int] = []
        for _, pids in neo_blocks:
            all_pids.extend(pids)
        meta = exp_v3._fetch_posts_meta(list(set(all_pids)), wl)
        need = [p for p in all_pids if p in meta and p not in exclude]
        emb = exp_v3._load_embeddings(need)
        snippets = _fetch_snippets(need)

        for sid, pids in neo_blocks:
            filt = [p for p in pids if p in meta and p not in exclude]
            if len(filt) < 3 or len(filt) > 12:
                continue
            filt.sort(key=lambda p: meta[p]["published_at"])
            ts_list = [meta[p]["published_at"] for p in filt]
            if min(ts_list) < MIN_PUBLISHED:
                continue
            if max(ts_list) < date_from or min(ts_list) > date_to:
                continue
            posts_2025 = sum(1 for t in ts_list if t >= date_from)
            if posts_2025 < max(2, len(filt) // 2):
                continue
            if any(meta[p]["text_len"] < 100 for p in filt):
                continue
            span_h = (max(ts_list) - min(ts_list)).total_seconds() / 3600.0
            if span_h < 12:
                continue
            bid = f"neo_dev_{sid}_{len(filt)}"
            if bid in seen:
                continue
            seen.add(bid)
            raw = {
                "candidate_block_id": bid,
                "source": "neo4j_storyline",
                "weak_storyline_id": sid,
                "post_ids": filt,
                "graph_cohesion_score": min(1.0, 0.25 + 0.07 * len(filt) + 0.06 * len({meta[p]['channel_id'] for p in filt})),
                "entity_overlap_proxy": "",
                "selection_reason": "neo_storyline_developmental_mining",
            }
            blocks.append(_enrich_block(raw, meta, emb, snippets))
    return blocks


def _mine_neo_near_duplicate(
    exclude: set[int],
    wl: set[str],
    date_from: datetime,
    date_to: datetime,
) -> list[dict[str, Any]]:
    """Short-span high-similarity blocks for diagnostic subset."""
    blocks: list[dict[str, Any]] = []
    seen: set[str] = set()
    for skip in range(0, 8000, 400):
        neo_blocks = exp_v3._neo_storyline_blocks(skip, 400)
        if not neo_blocks:
            break
        all_pids: list[int] = []
        for _, pids in neo_blocks:
            all_pids.extend(pids)
        meta = exp_v3._fetch_posts_meta(list(set(all_pids)), wl)
        need = [p for p in all_pids if p in meta and p not in exclude]
        emb = exp_v3._load_embeddings(need)
        snippets = _fetch_snippets(need)
        for sid, pids in neo_blocks:
            filt = [p for p in pids if p in meta and p not in exclude]
            if len(filt) < 3 or len(filt) > 8:
                continue
            filt.sort(key=lambda p: meta[p]["published_at"])
            ts_list = [meta[p]["published_at"] for p in filt]
            if max(ts_list) < date_from:
                continue
            span_h = (max(ts_list) - min(ts_list)).total_seconds() / 3600.0
            if span_h > 24:
                continue
            bid = f"neo_nd_{sid}_{len(filt)}"
            if bid in seen:
                continue
            seen.add(bid)
            raw = {
                "candidate_block_id": bid,
                "source": "neo4j_storyline",
                "weak_storyline_id": sid,
                "post_ids": filt,
                "graph_cohesion_score": 0.3,
                "entity_overlap_proxy": "",
                "selection_reason": "neo_near_duplicate_diagnostic",
            }
            bb = _enrich_block(raw, meta, emb, snippets)
            if (
                float(bb.get("near_duplicate_score", 0)) >= 0.45
                or (
                    float(bb.get("mean_pairwise_semantic", 0)) >= 0.82
                    and float(bb.get("high_semantic_pair_share", 0)) >= 0.55
                )
            ):
                bb["candidate_type"] = "near_duplicate_diagnostic"
                bb["near_duplicate_score"] = max(float(bb.get("near_duplicate_score", 0)), 0.55)
                blocks.append(bb)
    return blocks


def _mine_topic_window_developmental(
    exclude: set[int],
    wl_list: list[str],
    wl: set[str],
    date_from: datetime,
    date_to: datetime,
    topic: str,
    keywords: tuple[str, ...],
) -> list[dict[str, Any]]:
    """Group topic posts into 7-day windows; keep windows with span>=24h and size>=3."""
    pids = exp_v3._topic_search_posts(wl_list, exclude, keywords, date_from, date_to, 400)
    meta = exp_v3._fetch_posts_meta(pids, wl)
    by_week: dict[str, list[int]] = defaultdict(list)
    for pid, m in meta.items():
        wk = m["published_at"].strftime("%Y-W%W")
        by_week[wk].append(pid)
    blocks: list[dict[str, Any]] = []
    for wk, ps in by_week.items():
        if len(ps) < 3:
            continue
        ps = sorted(ps, key=lambda p: meta[p]["published_at"])
        ts_list = [meta[p]["published_at"] for p in ps]
        span_h = (max(ts_list) - min(ts_list)).total_seconds() / 3600.0
        if span_h < 24:
            continue
        if len(ps) > 8:
            ps = ps[:8]
        emb = exp_v3._load_embeddings(ps)
        snippets = _fetch_snippets(ps)
        bid = f"topic_dev_{topic}_{wk}"
        raw = {
            "candidate_block_id": bid,
            "source": "keyword_developmental",
            "weak_storyline_id": "",
            "post_ids": ps,
            "graph_cohesion_score": 0.2,
            "entity_overlap_proxy": "",
            "selection_reason": f"topic_window_{topic}",
        }
        bb = _enrich_block(raw, meta, emb, snippets)
        bb["topic_tags"] = topic if topic in bb.get("topic_tags", "") else f"{topic}," + bb.get("topic_tags", "other")
        if float(bb.get("developmental_score", 0)) >= 0.4 and float(bb.get("near_duplicate_score", 0)) < 0.55:
            blocks.append(bb)
    return blocks


def _mine_singleton_controls(
    exclude: set[int],
    wl_list: list[str],
    date_from: datetime,
    date_to: datetime,
    limit: int = 400,
) -> list[dict[str, Any]]:
    ch_arr = "ARRAY[" + ",".join("'" + c.replace("'", "''") + "'" for c in wl_list) + "]::text[]"
    excl = ""
    if exclude:
        excl = " AND id NOT IN (" + ",".join(str(x) for x in sorted(exclude)) + ")"
    sql = f"""
    SELECT coalesce(json_agg(post_id), '[]'::json)::text FROM (
      SELECT id AS post_id FROM public.posts
      WHERE timestamp >= '{date_from.isoformat()}' AND timestamp <= '{date_to.isoformat()}'
        AND channel_id = ANY({ch_arr}) AND length(coalesce(content,'')) >= 180 {excl}
      ORDER BY timestamp DESC LIMIT {limit}
    ) q;
    """
    pids = [int(x) for x in json.loads(exp_v3._psql_scalar(sql) or "[]")]
    meta = exp_v3._fetch_posts_meta(pids, set(wl_list))
    snippets = _fetch_snippets(list(meta.keys()))
    out: list[dict[str, Any]] = []
    for pid in pids:
        if pid not in meta:
            continue
        m = meta[pid]
        text = snippets.get(pid, "")
        tags = ",".join(exp_v3._topic_tags_from_text(text)) or "other"
        out.append(
            {
                "candidate_block_id": f"sgl_ctrl_{pid}",
                "source": "singleton_control",
                "weak_storyline_id": "",
                "post_ids": [pid],
                "size": 1,
                "story_size_bucket": "1",
                "channel_count": 1,
                "time_span_hours": 0.0,
                "time_span_bucket": "singleton",
                "temporal_phase_count": 1,
                "first_published_at": m["published_at"].isoformat(),
                "last_published_at": m["published_at"].isoformat(),
                "month_bucket": m["published_at"].strftime("%Y-%m"),
                "mean_pairwise_semantic": 0.0,
                "max_pairwise_semantic": 0.0,
                "high_semantic_pair_share": 0.0,
                "lexical_novelty_score": 0.0,
                "entity_novelty_score": 0.0,
                "graph_cohesion_score": 0.0,
                "entity_overlap_proxy": "",
                "developmental_score": 0.0,
                "near_duplicate_score": 0.0,
                "topic_tags": tags,
                "quality_flags": "singleton_control",
                "selection_reason": "singleton_control_stream",
                "candidate_type": "singleton_control",
            }
        )
    return out


def _mine_near_duplicate_diagnostic(
    pool: list[dict[str, Any]],
    limit: int = 12,
) -> list[dict[str, Any]]:
    cands = [
        b
        for b in pool
        if b.get("candidate_type") == "near_duplicate_diagnostic"
        or float(b.get("near_duplicate_score", 0)) >= 0.55
    ]
    cands.sort(key=lambda x: (-float(x.get("near_duplicate_score", 0)), -x["size"]))
    out: list[dict[str, Any]] = []
    for b in cands[:limit]:
        bb = dict(b)
        bb["candidate_type"] = "near_duplicate_diagnostic"
        bb["selection_reason_v3c"] = "near_duplicate_diagnostic_subset"
        out.append(bb)
    return out


def _mine_hard_negatives(
    dev_blocks: list[dict[str, Any]],
    all_blocks: list[dict[str, Any]],
    exclude: set[int],
    wl: set[str],
    max_blocks: int = 8,
) -> list[dict[str, Any]]:
    """Topic/time neighbors with different weak storyline."""
    chosen_sids = {b.get("weak_storyline_id") for b in dev_blocks}
    dev_topics: Counter[str] = Counter()
    dev_months: set[str] = set()
    for b in dev_blocks:
        for t in (b.get("topic_tags") or "").split(","):
            if t and t != "other":
                dev_topics[t] += 1
        dev_months.add(b.get("month_bucket", ""))

    cands: list[dict[str, Any]] = []
    for b in all_blocks:
        if b.get("weak_storyline_id") in chosen_sids:
            continue
        if b["size"] < 2:
            continue
        tags = set((b.get("topic_tags") or "").split(","))
        if not tags & set(dev_topics.keys()):
            continue
        if b.get("month_bucket") not in dev_months and dev_months:
            continue
        if float(b.get("near_duplicate_score", 0)) > 0.7:
            continue
        cands.append(b)
    cands.sort(key=lambda x: (-float(x.get("developmental_score", 0)), x["size"]))
    out: list[dict[str, Any]] = []
    used: set[int] = set()
    for b in cands:
        if any(p in exclude or p in used for p in b["post_ids"]):
            continue
        bb = dict(b)
        bb["candidate_type"] = "hard_negative"
        bb["selection_reason_v3c"] = "hard_negative_topic_neighbor"
        out.append(bb)
        used.update(b["post_ids"])
        if len(out) >= max_blocks:
            break
    return out


def _tags(b: dict[str, Any]) -> set[str]:
    return set((b.get("topic_tags") or "").split(","))


def _trump_iran_posts(block: dict[str, Any]) -> int:
    n = 0
    if "trump" in _tags(block) or "iran" in _tags(block):
        return block["size"]
    return n


def _span_diversity_bonus(span_h: float) -> float:
    if 24 <= span_h <= 72:
        return 0.35
    if 6 < span_h <= 24:
        return 0.3
    if 72 < span_h <= 168:
        return 0.25
    if span_h > 168:
        return 0.1
    return 0.0


def _select_review_set(
    developmental: list[dict[str, Any]],
    singletons: list[dict[str, Any]],
    near_dup: list[dict[str, Any]],
    hard_neg: list[dict[str, Any]],
    target: int = TARGET_POSTS,
) -> list[dict[str, Any]]:
    used: set[int] = set()
    chosen: list[dict[str, Any]] = []
    posts_by_type: Counter[str] = Counter()
    span_buckets: Counter[str] = Counter()
    trump_iran = 0
    quotas = {
        "developmental": int(0.62 * target),
        "singleton_control": int(0.22 * target),
        "near_duplicate_diagnostic": int(0.10 * target),
        "hard_negative": int(0.06 * target),
    }
    singleton_block_cap = 35
    singleton_blocks = 0

    def can_add(b: dict[str, Any], ctype: str) -> bool:
        pids = [p for p in b["post_ids"] if p not in used]
        if not pids:
            return False
        sz = len(pids)
        new_total = sum(posts_by_type.values()) + sz
        if new_total > TARGET_MAX:
            return False
        if len(chosen) >= MAX_BLOCKS:
            return False
        ti = len(pids) if ("trump" in _tags(b) or "iran" in _tags(b)) else 0
        if (trump_iran + ti) / max(new_total, 1) > TRUMP_IRAN_CAP + 1e-9:
            if not ("lunar" in _tags(b) or "dagestan_flood" in _tags(b)):
                return False
        if ctype == "near_duplicate_diagnostic":
            if posts_by_type[ctype] + sz > quotas[ctype] + 5:
                return False
        if ctype == "singleton_control":
            if singleton_blocks >= singleton_block_cap:
                return False
            if posts_by_type[ctype] + sz > quotas[ctype] + 8:
                return False
        if ctype == "hard_negative" and posts_by_type[ctype] + sz > quotas[ctype] + 8:
            return False
        return True

    def add_block(b: dict[str, Any], ctype: str, reason: str) -> bool:
        nonlocal trump_iran, singleton_blocks
        if not can_add(b, ctype):
            return False
        pids = [p for p in b["post_ids"] if p not in used]
        bb = dict(b)
        bb["post_ids"] = pids
        bb["size"] = len(pids)
        bb["story_size_bucket"] = _size_bucket(len(pids))
        bb["candidate_type"] = ctype
        bb["selection_reason_v3c"] = reason
        chosen.append(bb)
        used.update(pids)
        posts_by_type[ctype] += len(pids)
        if ctype == "developmental" and len(pids) >= 2:
            span_buckets[_span_bucket(float(bb.get("time_span_hours", 0)))] += len(pids)
        if "trump" in _tags(bb) or "iran" in _tags(bb):
            trump_iran += len(pids)
        if ctype == "singleton_control":
            singleton_blocks += 1
        return True

    dev_sorted = sorted(
        developmental,
        key=lambda x: (
            -("lunar" in _tags(x)) - ("dagestan_flood" in _tags(x)),
            -float(x.get("developmental_score", 0)),
            -_span_diversity_bonus(float(x.get("time_span_hours", 0))),
            -x["size"],
        ),
    )
    topics_seen: Counter[str] = Counter()

    for b in near_dup:
        if posts_by_type["near_duplicate_diagnostic"] >= quotas["near_duplicate_diagnostic"]:
            break
        add_block(b, "near_duplicate_diagnostic", "near_duplicate_diagnostic")

    for b in hard_neg:
        if posts_by_type["hard_negative"] >= quotas["hard_negative"]:
            break
        add_block(b, "hard_negative", "hard_negative")

    for b in dev_sorted:
        if posts_by_type["developmental"] >= quotas["developmental"]:
            break
        tg = _tags(b)
        topic_key = next((t for t in tg if t != "other"), "other")
        if topics_seen[topic_key] >= 6 and topic_key not in ("lunar", "dagestan_flood"):
            continue
        if float(b.get("developmental_score", 0)) < 0.42:
            continue
        if float(b.get("near_duplicate_score", 0)) >= 0.55:
            continue
        if add_block(b, "developmental", "developmental_primary"):
            topics_seen[topic_key] += 1

    for b in sorted(singletons, key=lambda x: ("lunar" in _tags(x), "dagestan_flood" in _tags(x)), reverse=True):
        if posts_by_type["singleton_control"] >= quotas["singleton_control"]:
            break
        add_block(b, "singleton_control", "singleton_control")

    while sum(posts_by_type.values()) < TARGET_MIN:
        if posts_by_type["developmental"] < quotas["developmental"] + 40:
            if any(add_block(b, "developmental", "developmental_fill") for b in dev_sorted):
                continue
        if posts_by_type["near_duplicate_diagnostic"] < quotas["near_duplicate_diagnostic"]:
            if any(add_block(b, "near_duplicate_diagnostic", "near_dup_fill") for b in near_dup):
                continue
        if posts_by_type["singleton_control"] < int(SING_POST_MIN * TARGET_MIN):
            if any(add_block(b, "singleton_control", "singleton_fill") for b in singletons):
                continue
        break

    return chosen


def _write_labeling_sheet_v3c(rows: list[dict[str, str]], csv_path: Path, xlsx_path: Path) -> None:
    csv_path.parent.mkdir(parents=True, exist_ok=True)
    with csv_path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=LABEL_COLS)
        w.writeheader()
        for r in rows:
            w.writerow({k: r.get(k, "") for k in LABEL_COLS})

    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    type_fills = {
        "developmental": PatternFill("solid", fgColor="C6EFCE"),
        "singleton_control": PatternFill("solid", fgColor="DDEBF7"),
        "near_duplicate_diagnostic": PatternFill("solid", fgColor="FCE4D6"),
        "hard_negative": PatternFill("solid", fgColor="E2EFDA"),
    }
    wb = openpyxl.Workbook()
    ins = wb.active
    ins.title = "instructions"
    ins["A1"] = "Manual gold expansion v3c — developmental stories"
    lines = [
        "gold_story_id — главное поле ручной разметки (оставьте пустым до разметки).",
        "suggested_story_id — подсказка (EXP3C_BLOCK_xxxx), не gold.",
        "candidate_type=developmental — развивающийся сюжет; проверьте, что посты не просто пересказы.",
        "Если developmental block содержит несколько сюжетов — split: разные gold_story_id по постам.",
        "Если блок — один развивающийся сюжет — один gold_story_id на все строки блока.",
        "near_duplicate_diagnostic — перепечатки; можно один gold_story_id, note=duplicate_chain при необходимости.",
        "hard_negative — обычно отдельный gold_story_id, если это другой сюжет.",
        "singleton_control — самостоятельная новость / контроль потока.",
        "temporal_phase — эвристическая фаза (phase_1_initial … phase_4_later).",
        "action=drop для мусора или нерелевантного поста.",
    ]
    for i, t in enumerate(lines, start=3):
        ins[f"A{i}"] = t

    ws = wb.create_sheet("review")
    header_fill = PatternFill("solid", fgColor="D9E1F2")
    for j, col in enumerate(LABEL_COLS, start=1):
        c = ws.cell(row=1, column=j, value=col)
        c.font = Font(bold=True)
        c.fill = header_fill
    prev_block = None
    for i, r in enumerate(rows, start=2):
        is_first = r["candidate_block_id"] != prev_block
        prev_block = r["candidate_block_id"]
        ctype = r.get("candidate_type", "")
        row_fill = type_fills.get(ctype)
        for j, col in enumerate(LABEL_COLS, start=1):
            val = r.get(col, "")
            cell = ws.cell(row=i, column=j, value=val)
            if col == "text":
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            if row_fill and (is_first or col in ("candidate_type", "candidate_block_id", "suggested_story_id")):
                cell.fill = row_fill
        ws.row_dimensions[i].height = 72 if r.get("text") else 22
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(LABEL_COLS))}{len(rows)+1}"
    widths = {
        "A": 8,
        "B": 14,
        "C": 16,
        "D": 18,
        "E": 20,
        "F": 5,
        "G": 7,
        "H": 14,
        "I": 22,
        "J": 18,
        "K": 14,
        "L": 14,
        "M": 22,
        "N": 20,
        "O": 12,
        "P": 10,
        "Q": 55,
    }
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    wb.save(xlsx_path)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--target-posts", type=int, default=TARGET_POSTS)
    args = ap.parse_args()

    root = _exp_root()
    out_dir = root / "sync_out/data/prepared/manual_gold_expansion_v3c"
    reports = root / "sync_out/reports"
    sync_in = root / "sync_in"
    for p in (out_dir, reports, sync_in):
        p.mkdir(parents=True, exist_ok=True)

    def _load_pids(path: Path) -> set[int]:
        if not path.is_file():
            return set()
        return {int(r["post_id"]) for r in csv.DictReader(path.open(encoding="utf-8"))}

    exclude = _load_pids(root / "sync_out/data/prepared/combined_manual_gold_v2/combined_manual_gold_labels_v2.csv")
    exclude |= _load_pids(root / "sync_out/data/prepared/manual_gold_expansion_v3/selected_posts_metadata_v3.csv")
    exclude |= _load_pids(root / "sync_out/data/prepared/manual_gold_expansion_v3b/selected_posts_metadata_v3b.csv")
    v3b_pids = _load_pids(root / "sync_out/data/prepared/manual_gold_expansion_v3b/selected_posts_metadata_v3b.csv")

    wl_path = root / "sync_out/data/prepared/dataset_candidate_mining_whitelist_v1/channel_whitelist_v1.csv"
    wl: dict[str, dict[str, str]] = {}
    for r in csv.DictReader(wl_path.open(encoding="utf-8")):
        wl[r["channel_id"].strip()] = r
    wl_ids = set(wl.keys())

    max_ts = exp_v3._parse_ts(exp_v3._psql_scalar("SELECT max(timestamp)::text FROM public.posts WHERE timestamp >= '2024-01-01';"))
    assert max_ts
    date_to = max_ts
    date_from = max(PRIMARY_START, MIN_PUBLISHED)

    neo_pool = _mine_neo_developmental(exclude, wl_ids, date_from, date_to)
    topic_dev = _mine_topic_window_developmental(
        exclude, sorted(wl_ids), wl_ids, date_from, date_to, "lunar", exp_v3.LUNAR_KW
    )
    topic_dev += _mine_topic_window_developmental(
        exclude, sorted(wl_ids), wl_ids, date_from, date_to, "dagestan_flood", exp_v3.DAGESTAN_KW
    )
    seen_p: set[frozenset[int]] = set()
    merged: list[dict[str, Any]] = []
    for b in neo_pool + topic_dev:
        key = frozenset(b["post_ids"])
        if key in seen_p:
            continue
        seen_p.add(key)
        merged.append(b)
    neo_pool = merged
    developmental = [b for b in neo_pool if b.get("candidate_type") != "reject" and float(b.get("developmental_score", 0)) >= 0.4]
    near_pool = [b for b in neo_pool if float(b.get("near_duplicate_score", 0)) >= 0.5]
    near_pool += _mine_neo_near_duplicate(exclude, wl_ids, date_from, date_to)

    singletons = _mine_singleton_controls(exclude, sorted(wl_ids), date_from, date_to)

    # write candidate pool
    with (out_dir / "candidate_blocks_v3c.csv").open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=CAND_FIELDS)
        w.writeheader()
        for b in neo_pool + singletons:
            row = {k: b.get(k, "") for k in CAND_FIELDS}
            row["post_ids"] = ";".join(str(p) for p in b["post_ids"])
            w.writerow(row)

    dev_spans = [float(b["time_span_hours"]) for b in developmental if b["size"] >= 3]
    inv_md = (
        "# Candidate inventory v3c (developmental)\n\n"
        f"- Neo blocks mined (span≥12h, size≥3): **{len(neo_pool)}**\n"
        f"- Developmental candidates (score≥0.4): **{len(developmental)}**\n"
        f"- Near-duplicate flagged: **{sum(1 for b in neo_pool if float(b.get('near_duplicate_score',0))>=0.5)}**\n"
        f"- Median span (developmental): **{float(np.median(dev_spans)) if dev_spans else 0:.1f}** h\n"
        f"- Singleton controls available: **{len(singletons)}**\n"
        f"- Date window: {date_from.date()} … {date_to.date()}\n"
    )
    reports.joinpath("manual_gold_expansion_candidate_inventory_v3c.md").write_text(inv_md, encoding="utf-8")

    near_diag = _mine_near_duplicate_diagnostic(near_pool, limit=18)
    chosen = _select_review_set(
        developmental,
        singletons,
        near_diag,
        _mine_hard_negatives(developmental, neo_pool, exclude, wl_ids),
        target=args.target_posts,
    )

    reused_v3b = {p for b in chosen for p in b["post_ids"]} & v3b_pids

    label_rows: list[dict[str, str]] = []
    meta_rows: list[dict[str, str]] = []
    block_rows: list[dict[str, Any]] = []
    pids_all = [p for b in chosen for p in b["post_ids"]]
    arr = "ARRAY[" + ",".join(str(x) for x in pids_all) + "]::bigint[]"
    pg_rows = exp_v3._psql_json(
        f"""
        SELECT coalesce(json_agg(row_to_json(t) ORDER BY t.post_id), '[]'::json)::text FROM (
          SELECT p.id AS post_id, p.timestamp AS published_at, p.channel_id,
                 coalesce(c.username, '') AS username, coalesce(c.name, '') AS channel_title,
                 coalesce(p.content, '') AS content
          FROM public.posts p
          LEFT JOIN communities c ON c.chat_id::text = p.channel_id
          WHERE p.id = ANY({arr})
        ) t;
        """
    )
    pg_map = {int(r["post_id"]): r for r in pg_rows}

    months: Counter[str] = Counter()
    old2023 = 0
    min_dt = max_dt = None
    posts_by_type: Counter[str] = Counter()
    dev_spans_sel: list[float] = []
    phase_counts: list[int] = []

    for bi, b in enumerate(chosen, start=1):
        sug = f"EXP3C_BLOCK_{bi:04d}"
        ctype = b.get("candidate_type", "developmental")
        block_rows.append(
            {
                "candidate_block_id": b["candidate_block_id"],
                "suggested_story_id": sug,
                "candidate_type": ctype,
                "n_posts": len(b["post_ids"]),
                "time_span_hours": b.get("time_span_hours", ""),
                "temporal_phase_count": b.get("temporal_phase_count", ""),
                "developmental_score": b.get("developmental_score", ""),
                "near_duplicate_score": b.get("near_duplicate_score", ""),
                "topic_tags": b.get("topic_tags", ""),
                "selection_reason_v3c": b.get("selection_reason_v3c", ""),
            }
        )
        if ctype == "developmental":
            dev_spans_sel.append(float(b.get("time_span_hours") or 0))
            phase_counts.append(int(b.get("temporal_phase_count") or 0))
        t0 = None
        sorted_pids = sorted(
            b["post_ids"],
            key=lambda p: exp_v3._parse_ts(str(pg_map[p]["published_at"])) or MIN_PUBLISHED,
        )
        for ri, pid in enumerate(sorted_pids):
            pg = pg_map.get(pid)
            if not pg:
                continue
            ts = exp_v3._parse_ts(str(pg["published_at"]))
            if not ts:
                continue
            if t0 is None:
                t0 = ts
            hours = (ts - t0).total_seconds() / 3600.0
            phase = _phase_label(hours) if ctype == "developmental" else "singleton"
            if ts < MIN_PUBLISHED:
                old2023 += 1
            min_dt = ts if min_dt is None else min(min_dt, ts)
            max_dt = ts if max_dt is None else max(max_dt, ts)
            months[ts.strftime("%Y-%m")] += 1
            posts_by_type[ctype] += 1
            uname = (pg.get("username") or "").strip()
            ch_name = uname if uname.startswith("@") else (f"@{uname}" if uname else str(pg.get("channel_title") or ""))
            text = exp_v3._normalize_text(str(pg.get("content") or ""))
            tags = b.get("topic_tags") or ",".join(exp_v3._topic_tags_from_text(text))
            meta_rows.append(
                {
                    "post_id": pid,
                    "published_at": ts.isoformat(),
                    "channel_id": str(pg.get("channel_id") or ""),
                    "channel_name": ch_name,
                    "candidate_block_id": b["candidate_block_id"],
                    "suggested_story_id": sug,
                    "candidate_type": ctype,
                    "block_size": len(b["post_ids"]),
                    "story_size_bucket": b.get("story_size_bucket", ""),
                    "topic_tags": tags,
                }
            )
            label_rows.append(
                {
                    "action": "",
                    "gold_story_id": "",
                    "note": "",
                    "suggested_story_id": sug,
                    "candidate_block_id": b["candidate_block_id"],
                    "row_in_block": ri + 1,
                    "block_size": len(b["post_ids"]),
                    "weak_storyline_id": b.get("weak_storyline_id", ""),
                    "candidate_type": ctype,
                    "temporal_phase": phase,
                    "topic_tags": tags,
                    "selection_reason": b.get("selection_reason_v3c", b.get("selection_reason", "")),
                    "published_at": ts.strftime("%Y-%m-%d %H:%M:%S%z"),
                    "channel_name": ch_name,
                    "post_id": pid,
                    "text": text,
                }
            )

    total = len(label_rows)
    bucket_posts = Counter()
    for b in chosen:
        bucket_posts[_size_bucket(len(b["post_ids"]))] += len(b["post_ids"])
    trump_posts = sum(
        len(b["post_ids"]) for b in chosen if "trump" in _tags(b) or "iran" in _tags(b)
    )
    lunar_posts = sum(len(b["post_ids"]) for b in chosen if "lunar" in _tags(b))
    dag_posts = sum(len(b["post_ids"]) for b in chosen if "dagestan_flood" in _tags(b))

    manifest = {
        "dataset": "manual_gold_expansion_v3c",
        "target_posts": args.target_posts,
        "selected_posts": total,
        "selected_blocks": len(chosen),
        "overlap_combined_manual_gold_v2": 0,
        "overlap_v3_count": len({p for b in chosen for p in b["post_ids"]} & _load_pids(
            root / "sync_out/data/prepared/manual_gold_expansion_v3/selected_posts_metadata_v3.csv"
        )),
        "overlap_v3b_count": len(reused_v3b),
        "reused_from_v3b_count": len(reused_v3b),
        "candidate_type_post_counts": dict(posts_by_type),
        "developmental_post_rate": round(posts_by_type["developmental"] / max(total, 1), 4),
        "singleton_control_post_rate": round(posts_by_type["singleton_control"] / max(total, 1), 4),
        "near_duplicate_diagnostic_post_rate": round(
            posts_by_type["near_duplicate_diagnostic"] / max(total, 1), 4
        ),
        "hard_negative_post_rate": round(posts_by_type["hard_negative"] / max(total, 1), 4),
        "story_size_bucket_distribution": dict(bucket_posts),
        "time_span_bucket_posts": dict(
            Counter(
                _span_bucket(float(b.get("time_span_hours", 0)))
                for b in chosen
                if b.get("candidate_type") == "developmental"
                for _ in b["post_ids"]
            )
        ),
        "developmental_median_time_span_hours": float(np.median(dev_spans_sel)) if dev_spans_sel else 0.0,
        "developmental_mean_time_span_hours": float(np.mean(dev_spans_sel)) if dev_spans_sel else 0.0,
        "developmental_blocks_phase_ge2_share": round(
            sum(1 for c in phase_counts if c >= 2) / max(len(phase_counts), 1), 4
        ),
        "trump_iran_rate": round(trump_posts / max(total, 1), 4),
        "lunar_count": lunar_posts,
        "dagestan_count": dag_posts,
        "published_at_min": min_dt.isoformat() if min_dt else "",
        "published_at_max": max_dt.isoformat() if max_dt else "",
        "month_distribution": dict(months),
        "2023_or_older_count": old2023,
        "whitelist_source": str(wl_path),
        "whitelist_only": True,
        "candidate_pool_blocks": len(neo_pool),
        "date_window": {"from": date_from.isoformat(), "to": date_to.isoformat()},
    }

    (out_dir / "selection_manifest_v3c.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    with (out_dir / "selected_blocks_v3c.csv").open("w", encoding="utf-8", newline="") as f:
        if block_rows:
            w = csv.DictWriter(f, fieldnames=list(block_rows[0].keys()))
            w.writeheader()
            w.writerows(block_rows)
    with (out_dir / "selected_posts_metadata_v3c.csv").open("w", encoding="utf-8", newline="") as f:
        if meta_rows:
            w = csv.DictWriter(f, fieldnames=list(meta_rows[0].keys()))
            w.writeheader()
            w.writerows(meta_rows)

    reports.joinpath("manual_gold_expansion_selection_v3c.md").write_text(
        "# Selection v3c (developmental stories)\n\n"
        f"- Posts: **{total}**, blocks: **{len(chosen)}**\n"
        f"- By type: `{dict(posts_by_type)}`\n"
        f"- Developmental rate: **{manifest['developmental_post_rate']:.1%}**\n"
        f"- Near-dup diagnostic: **{manifest['near_duplicate_diagnostic_post_rate']:.1%}**\n"
        f"- Median developmental span: **{manifest['developmental_median_time_span_hours']:.1f}** h\n"
        f"- Overlap v3b reused: **{len(reused_v3b)}** posts\n"
        f"- Trump/Iran: **{manifest['trump_iran_rate']:.1%}**\n",
        encoding="utf-8",
    )

    _write_labeling_sheet_v3c(
        label_rows,
        sync_in / "grouped_manual_gold_expansion_v3c_developmental_labeling.csv",
        sync_in / "grouped_manual_gold_expansion_v3c_developmental_labeling.xlsx",
    )
    print(json.dumps(manifest, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
